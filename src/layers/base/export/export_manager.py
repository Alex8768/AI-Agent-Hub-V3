from __future__ import annotations

import base64
from dataclasses import dataclass
from pathlib import Path
from src.core.config import settings
from uuid import uuid4

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from src.infrastructure.database.models import DocumentRecord


@dataclass(frozen=True)
class ExportArtifact:
    export_id: str
    format: str
    output_path: str          # relative path under ./data
    size_bytes: int
    download_url: str


class ExportManager:
    """
    Base ExportManager:
    - takes content OR document_id
    - writes artifact into settings.data_dir / "exports" / <export_id>.<ext>
    - returns ExportArtifact (download via /api/v1/export/{export_id}/download)
    """

    def __init__(self, exports_root: str | None = None):
        exports_root = exports_root or str(settings.data_dir / "exports")
        self.root = Path(exports_root)
        self.root.mkdir(parents=True, exist_ok=True)

    async def _load_text_from_document(self, db: AsyncSession, document_id: str) -> tuple[str, str]:
        """
        Returns: (text, filename)
        Base: assumes stored file is text-ish (utf-8 decode with ignore).
        """
        stmt = select(DocumentRecord).where(DocumentRecord.id == document_id)
        res = await db.execute(stmt)
        rec = res.scalar_one_or_none()
        if rec is None:
            raise ValueError(f"Document not found: {document_id}")

        storage_path = Path(settings.data_dir) / rec.storage_key
        if not storage_path.exists():
            raise FileNotFoundError(f"Stored file not found: {storage_path}")

        data = storage_path.read_bytes()
        text = data.decode("utf-8", errors="ignore")
        return text, rec.filename

    def _write_md_or_txt(self, out_path: Path, text: str) -> None:
        out_path.write_text(text, encoding="utf-8")

    def _write_pdf(self, out_path: Path, text: str, title: str | None = None) -> None:
        # reportlab is installed
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import A4

        c = canvas.Canvas(str(out_path), pagesize=A4)
        width, height = A4

        y = height - 50
        if title:
            c.setFont("Helvetica-Bold", 14)
            c.drawString(50, y, title)
            y -= 30

        c.setFont("Helvetica", 11)
        # simple wrapping
        max_chars = 95
        for para in text.splitlines() or [""]:
            line = para
            while len(line) > max_chars:
                c.drawString(50, y, line[:max_chars])
                y -= 15
                line = line[max_chars:]
                if y < 60:
                    c.showPage()
                    c.setFont("Helvetica", 11)
                    y = height - 50
            c.drawString(50, y, line)
            y -= 15
            if y < 60:
                c.showPage()
                c.setFont("Helvetica", 11)
                y = height - 50

        c.save()

    def _write_docx(self, out_path: Path, text: str, title: str | None = None) -> None:
        # python-docx is installed
        from docx import Document

        doc = Document()
        if title:
            doc.add_heading(title, level=1)

        # keep newlines as paragraphs
        for line in text.splitlines() or [""]:
            doc.add_paragraph(line)

        doc.save(str(out_path))

    async def export(
        self,
        db: AsyncSession,
        *,
        format: str,
        content: str | None = None,
        document_id: str | None = None,
    ) -> ExportArtifact:
        fmt = (format or "md").lower().strip()
        if fmt not in {"md", "txt", "pdf", "docx"}:
            raise ValueError(f"Unsupported export format: {fmt}")

        if content is None:
            if not document_id:
                raise ValueError("Either content or document_id must be provided")
            content, filename = await self._load_text_from_document(db, document_id)
            title = filename
        else:
            title = None

        export_id = uuid4().hex
        ext = "md" if fmt == "md" else fmt
        out_file = self.root / f"{export_id}.{ext}"

        if fmt in {"md", "txt"}:
            self._write_md_or_txt(out_file, content)
        elif fmt == "pdf":
            self._write_pdf(out_file, content, title=title)
        elif fmt == "docx":
            self._write_docx(out_file, content, title=title)

        rel = str(out_file.relative_to(Path(settings.data_dir)))
        size = out_file.stat().st_size
        download_url = f"/api/v1/export/{export_id}/download"

        return ExportArtifact(
            export_id=export_id,
            format=fmt,
            output_path=rel,
            size_bytes=size,
            download_url=download_url,
        )

    def find_export_file(self, export_id: str) -> Path | None:
        matches = list(self.root.glob(f"{export_id}.*"))
        return matches[0] if matches else None
